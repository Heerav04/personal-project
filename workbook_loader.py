"""
Workbook loader and inspector module.
Provides robust, case-insensitive, whitespace-tolerant discovery of sheets and headers
using configurable candidate names, with inspect capabilities.
"""

from __future__ import annotations
import logging
import re
from pathlib import Path
from typing import Any, Dict, Generator, List, Optional, Tuple
import openpyxl

from config import Config

logger = logging.getLogger(__name__)


def normalize_header(header: Any) -> str:
    """Normalize a header string for tolerant comparison."""
    if header is None:
        return ""
    # Convert to lowercase and strip all non-alphanumeric chars
    return re.sub(r"[^a-z0-9]", "", str(header).lower())


class WorkbookInspector:
    """Inspects an Excel workbook structure and reports sheet names and headers."""

    @staticmethod
    def inspect(file_path: str | Path) -> Dict[str, Dict[str, Any]]:
        wb_path = Path(file_path)
        if not wb_path.exists():
            raise FileNotFoundError(f"Workbook not found: {wb_path}")

        wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
        report: Dict[str, Dict[str, Any]] = {}

        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                headers: List[str] = []
                sample_rows: List[List[Any]] = []
                row_count = 0

                for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                    if not any(cell is not None and str(cell).strip() != "" for cell in row):
                        continue
                    row_count += 1
                    if not headers:
                        headers = [str(c).strip() if c is not None else f"Column_{i+1}" for i, c in enumerate(row)]
                    elif len(sample_rows) < 2:
                        sample_rows.append([str(c).strip() if c is not None else "" for c in row[:len(headers)]])

                report[sheet_name] = {
                    "row_count": row_count,
                    "headers": headers,
                    "sample_rows": sample_rows
                }
        finally:
            wb.close()

        return report

    @staticmethod
    def print_inspection(file_path: str | Path) -> None:
        report = WorkbookInspector.inspect(file_path)
        print(f"\n=======================================================")
        print(f" WORKBOOK INSPECTION REPORT: {Path(file_path).name}")
        print(f"=======================================================\n")
        for sheet_name, info in report.items():
            print(f"Sheet: '{sheet_name}' (Total Rows: {info['row_count']})")
            print(f"  Detected Headers ({len(info['headers'])}):")
            for idx, h in enumerate(info["headers"], 1):
                print(f"    [{idx}] {h}")
            if info["sample_rows"]:
                print(f"  Sample Data (Row 1):")
                sample = info["sample_rows"][0]
                for h, val in zip(info["headers"][:8], sample[:8]):
                    print(f"      {h}: {val}")
            print("-" * 55)


class WorkbookLoader:
    """Loads and navigates sheets and columns using configurable rules."""

    def __init__(self, file_path: str | Path, config: Config):
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"Excel workbook not found: {self.file_path}")
        self.config = config
        self.wb = openpyxl.load_workbook(self.file_path, data_only=True)

    def close(self) -> None:
        if self.wb:
            self.wb.close()

    def __enter__(self) -> WorkbookLoader:
        return self

    def __exit__(self, exc_type, exc_val, exc_tb) -> None:
        self.close()

    def resolve_sheet_name(self, sheet_type: str) -> str:
        """
        Locates the actual sheet name from candidate names in config.
        """
        candidates = self.config.get_sheet_candidates(sheet_type)
        actual_sheets = self.wb.sheetnames

        # 1. Exact match
        for cand in candidates:
            if cand in actual_sheets:
                return cand

        # 2. Case-insensitive & normalized match
        norm_actual = {normalize_header(s): s for s in actual_sheets}
        for cand in candidates:
            norm_cand = normalize_header(cand)
            if norm_cand in norm_actual:
                return norm_actual[norm_cand]

        # 3. Substring match
        for cand in candidates:
            norm_cand = normalize_header(cand)
            for norm_act, act_name in norm_actual.items():
                if norm_cand and (norm_cand in norm_act or norm_act in norm_cand):
                    return act_name

        raise ValueError(
            f"Could not find worksheet for '{sheet_type}'. Looked for candidates: {candidates}. "
            f"Available sheets in workbook: {actual_sheets}"
        )

    def map_sheet_columns(
        self,
        sheet_name: str,
        column_roles: Dict[str, List[str]]
    ) -> Tuple[Dict[str, int], List[str]]:
        """
        Maps logical column roles to 0-based column indices.
        Returns (mapping_dict, raw_header_list).
        """
        ws = self.wb[sheet_name]
        raw_headers: List[str] = []
        header_row_idx = None

        # Find the header row (first non-empty row)
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            non_empty = [c for c in row if c is not None and str(c).strip() != ""]
            if non_empty:
                header_row_idx = row_idx
                raw_headers = [str(c).strip() if c is not None else "" for c in row]
                break

        if header_row_idx is None:
            raise ValueError(f"Worksheet '{sheet_name}' is completely empty.")

        norm_headers = [normalize_header(h) for h in raw_headers]
        role_to_idx: Dict[str, int] = {}

        for role, candidates in column_roles.items():
            matched_idx = None

            # Try exact match
            for cand in candidates:
                for idx, h in enumerate(raw_headers):
                    if cand.lower() == h.lower():
                        matched_idx = idx
                        break
                if matched_idx is not None:
                    break

            # Try normalized match
            if matched_idx is None:
                for cand in candidates:
                    norm_c = normalize_header(cand)
                    for idx, nh in enumerate(norm_headers):
                        if norm_c and norm_c == nh:
                            matched_idx = idx
                            break
                    if matched_idx is not None:
                        break

            # Try substring match
            if matched_idx is None:
                for cand in candidates:
                    norm_c = normalize_header(cand)
                    if len(norm_c) >= 3:
                        for idx, nh in enumerate(norm_headers):
                            if norm_c in nh or nh in norm_c:
                                matched_idx = idx
                                break
                    if matched_idx is not None:
                        break

            if matched_idx is not None:
                role_to_idx[role] = matched_idx
            else:
                logger.warning(
                    f"Optional column role '{role}' could not be matched in sheet '{sheet_name}'. "
                    f"Looked for: {candidates}. Available headers: {raw_headers}"
                )

        return role_to_idx, raw_headers

    def read_rows_as_dicts(
        self,
        sheet_name: str,
        role_to_idx: Dict[str, int],
        skip_header: bool = True
    ) -> Generator[Dict[str, Any], None, None]:
        """
        Iterates over rows in sheet and yields mapped dicts.
        """
        ws = self.wb[sheet_name]
        first_row_seen = False

        for row in ws.iter_rows(values_only=True):
            if not any(cell is not None and str(cell).strip() != "" for cell in row):
                continue

            if not first_row_seen and skip_header:
                first_row_seen = True
                continue

            record: Dict[str, Any] = {}
            for role, idx in role_to_idx.items():
                if idx < len(row):
                    val = row[idx]
                    if val is not None:
                        val_str = str(val).strip()
                        record[role] = val_str if val_str != "" else None
                    else:
                        record[role] = None
                else:
                    record[role] = None

            yield record
